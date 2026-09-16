"""Build Word documents for the reviewer-facing application downloads."""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "reviewer_documents"
OUT.mkdir(parents=True, exist_ok=True)

INK = "24213D"
MUTED = "645D70"
PURPLE = "7152D6"
NAVY = "32304A"
LAVENDER = "EEE8FF"
PALE = "F7F5FB"
LINE = "D9D5E2"
WHITE = "FFFFFF"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=140, bottom=90, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def page_field(paragraph):
    paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, before, after in (("Title", 27, 0, 10), ("Subtitle", 12.5, 0, 14), ("Heading 1", 17.5, 13, 6), ("Heading 2", 13.5, 10, 4)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "SECURELAUNCH AI  |  PORTFOLIO EVIDENCE"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(PURPLE)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Eniola Durojaiye  |  AI security case study  |  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    page_field(footer)


def add_title(doc, title, subtitle, created, revised="September 16, 2026"):
    doc.add_paragraph(title, style="Title")
    doc.add_paragraph(subtitle, style="Subtitle")
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(5.15)
    rows = [("Author", "Eniola Durojaiye"), ("Analysis dates", created), ("Reviewer edition", revised)]
    for i, (label, value) in enumerate(rows):
        for cell in table.rows[i].cells:
            set_cell_borders(cell, color=WHITE, size="0")
            set_cell_margins(cell, top=70, bottom=70)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(table.cell(i, 0), LAVENDER)
        p = table.cell(i, 0).paragraphs[0]
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(INK)
        table.cell(i, 1).text = value
    doc.add_paragraph()


def add_scope_note(doc):
    p = doc.add_paragraph()
    r = p.add_run("Evidence boundary  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(PURPLE)
    p.add_run(
        "I created LindenArc as a realistic business scenario and used synthetic records so I could test security controls without using real customer information. "
        "The Python gateway, automated tests, risk analysis, Amazon Bedrock Guardrail configuration, and recorded AWS test results are real project work."
    )


def add_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3.5)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_fill(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            set_cell_fill(cell, PALE if i % 2 else WHITE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    run.font.size = Pt(9.4)
    if widths:
        for row in table.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Inches(width)
    doc.add_paragraph()
    return table


def add_code(doc, code):
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line or " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(8.8)
        run.font.color.rgb = RGBColor.from_string(INK)


def finish(doc, filename):
    path = OUT / filename
    doc.core_properties.author = "Eniola Durojaiye"
    doc.core_properties.title = filename.removesuffix(".docx").replace("-", " ")
    doc.core_properties.subject = "SecureLaunch AI portfolio evidence"
    doc.save(path)
    return path


def scope_document():
    doc = Document(); style_doc(doc)
    add_title(doc, "SecureLaunch AI Project Scope", "Business scenario, security question, and assessment boundaries", "September 4-7, 2026")
    add_scope_note(doc)
    doc.add_heading("Purpose and conclusion", level=1)
    add_paragraph(doc, "I created this case study to examine how a financial software company could use an external AI service without sending unnecessary customer information outside the organization. My conclusion was that the AI route needed a security gateway before any external API request, a separate return-path validation step, and human review of every result.")
    doc.add_heading("The business scenario", level=1)
    add_paragraph(doc, "LindenArc is a fictional business-to-business financial software company that supports invoice, expense, approval, and payment workflows. I proposed Signal as an AI-assisted support feature that could summarize a ticket and identify an approved troubleshooting category. The external AI provider would supply a draft summary, but LindenArc would retain control over the customer record, the security policy, and the final support action.")
    doc.add_heading("The security question", level=1)
    add_paragraph(doc, "The central question was: how can support staff gain useful context from AI without expanding the exposure of customer information or granting the model business authority?")
    add_bullets(doc, [
        "Inspect and minimize a support ticket before constructing any external request.",
        "Mask approved lower-risk identifiers while preserving useful troubleshooting context.",
        "Block restricted identity, financial, credential, and manipulation cases.",
        "Validate the returned content before displaying it to a support agent.",
        "Keep the AI outside payment, account-change, communication, and publishing authority.",
    ])
    doc.add_heading("What Signal can and cannot do", level=1)
    add_table(doc, ["Permitted capability", "Required boundary"], [
        ("Summarize the support problem", "Only minimized and approved text may enter the external request"),
        ("Suggest an approved issue category", "The response must match a strict schema and allowed category list"),
        ("Help an agent find a playbook", "LindenArc retrieves the human-authored playbook after response validation"),
        ("Execute an account or payment action", "Prohibited. The model receives no tool, payment, or account-change authority"),
    ], [2.4, 4.3])
    doc.add_heading("My role and deliverables", level=1)
    add_paragraph(doc, "I acted as the security analyst for the scenario. I defined the system boundaries, mapped the data, completed the threat and risk analysis, developed and tested the Python security gateway, and independently configured and validated Amazon Bedrock Guardrails. I kept findings from the local prototype separate from the AWS evidence so each result remained accurate.")
    add_table(doc, ["Deliverable", "What it contributes"], [
        ("Scenario charter", "Defines the business need, trust boundary, permitted use, and exclusions"),
        ("Data inventory", "Identifies what information exists, where it moves, and what may reach AI"),
        ("PASTA threat model", "Connects the use case to threats, weaknesses, scenarios, and prioritized risks"),
        ("Python gateway", "Demonstrates minimization, masking, blocking, final inspection, and safe fallback"),
        ("AWS Guardrail lab", "Provides independent cloud evidence for selected data and prompt-attack cases"),
    ], [2.2, 4.5])
    doc.add_heading("Assessment limits", level=1)
    add_paragraph(doc, "The scenario does not claim a production deployment, universal detection accuracy, legal compliance, or approval to launch. The synthetic tests demonstrate specific control behavior. Production tenant isolation, supplier terms, access controls, retention, deletion, monitoring, resilience, and real-model quality would require separate evidence.")
    return finish(doc, "01-project-scope-and-security-boundary.docx")


def analysis_document():
    doc = Document(); style_doc(doc)
    add_title(doc, "SecureLaunch AI Data and Risk Analysis", "Data lifecycle, threat modeling, and prioritized security concerns", "September 5-9, 2026")
    add_scope_note(doc)
    doc.add_heading("Purpose and conclusion", level=1)
    add_paragraph(doc, "I traced the support-ticket lifecycle before selecting controls. This analysis showed that the highest-value portfolio question was not whether AI could summarize a ticket. It was whether LindenArc could prevent unnecessary or restricted information from entering the AI route and could stop unsafe output from influencing an agent.")
    doc.add_heading("How I followed the data", level=1)
    add_numbered(doc, [
        "The customer submits an original support ticket inside LindenArc.",
        "The gateway removes fields that the external AI does not need, including tenant metadata, contact fields, account references, and attachments.",
        "The gateway inspects the permitted text, masks selected identifiers, and blocks restricted information or manipulation attempts.",
        "The exact outgoing JSON receives a second inspection before a provider call is possible.",
        "The returned content is treated as untrusted, validated, and paired with a human-authored playbook for agent review.",
    ])
    doc.add_heading("Data decisions", level=1)
    add_table(doc, ["Information", "Classification and decision", "External AI treatment"], [
        ("Original support ticket", "Confidential by default; may contain restricted content", "Send minimized text only after inspection"),
        ("Attachments", "Confidential or restricted according to content", "Do not send"),
        ("Working AI payload", "Confidential; restricted if a detector misses content", "Approved AI input after masking or blocking"),
        ("Draft summary and category", "Confidential derived information", "Validate before display or storage"),
        ("Credentials and payment data", "Restricted", "Block the AI route"),
        ("Synthetic evaluation records", "Internal during development; public only after review", "Use for local and AWS testing without real customer data"),
    ], [1.65, 2.55, 2.55])
    doc.add_heading("How I used PASTA", level=1)
    add_paragraph(doc, "I applied all seven PASTA stages so the threat model would begin with business objectives and end with prioritized risk. I defined objectives, established scope, decomposed the system, analyzed threats, identified candidate weaknesses, modeled attack and failure scenarios, and prioritized risk. I recorded assumptions and evidence gaps rather than presenting design ideas as verified controls.")
    doc.add_heading("Priority risks", level=1)
    add_table(doc, ["Risk", "Why it matters", "Primary response"], [
        ("Sensitive-data disclosure", "Free-form tickets can contain personal, financial, credential, or proprietary information", "Minimize fields, inspect content, mask or block, and reinspect final JSON"),
        ("Misleading AI output", "A plausible response can still be wrong or manipulated", "Validate structure and content, restrict categories, and retain human review"),
        ("Provider dependency", "An external supplier can fail, change behavior, or expose data", "Minimize data, validate responses, use fallback, and require provider evidence"),
        ("Cross-tenant access", "A multi-tenant application can expose another customer's records if authorization is weak", "Require server-side tenant context, object authorization, and negative tests"),
        ("Privilege or change compromise", "A bad change can weaken several controls at once", "Use least privilege, reviewed changes, testing, monitoring, and rollback"),
    ], [1.6, 2.9, 2.25])
    doc.add_heading("What the portfolio verifies", level=1)
    add_bullets(doc, [
        "The Python implementation can remove fields, detect selected patterns, mask approved identifiers, and suppress provider invocation.",
        "The automated tests can verify fixed local scenarios, safe failure, and evidence association.",
        "The AWS lab can verify selected Guardrail behavior using synthetic text without a foundation-model call.",
    ])
    doc.add_heading("What remains unresolved", level=1)
    add_paragraph(doc, "The project does not verify production tenant isolation, real provider data practices, complete detector coverage, production IAM, operational monitoring, deletion, recovery, or sustainable human review. I preserved these as remaining risks because a functioning filter is one security layer, not a complete security program.")
    return finish(doc, "02-data-lifecycle-and-risk-analysis.docx")


def build_document():
    doc = Document(); style_doc(doc)
    add_title(doc, "SecureLaunch AI Python Security Gateway", "Control design, implementation, and technical behavior", "September 11, 2026")
    add_scope_note(doc)
    doc.add_heading("Purpose and conclusion", level=1)
    add_paragraph(doc, "I built a pre-provider Python gateway that turns the security policy into executable decisions. The gateway constructs a limited payload only after inspection. If a blocking condition appears, it returns a decision without creating a provider payload, so the external AI route cannot run.")
    doc.add_heading("Control sequence", level=1)
    add_numbered(doc, [
        "Validate the structure of the synthetic support ticket.",
        "Check the emergency kill switch.",
        "Exclude tenant metadata, separate contact fields, account references, and attachments from the provider payload.",
        "Enforce the ticket-size boundary and inspect normalized text.",
        "Mask approved identifiers or block restricted information and prompt manipulation.",
        "Construct a strict JSON payload containing only approved fields.",
        "Inspect the exact serialized JSON again before the local provider adapter can run.",
        "Validate the simulated response and retrieve a human-authored playbook for agent review.",
    ])
    doc.add_heading("Policy decisions", level=1)
    add_table(doc, ["Decision", "Examples", "Behavior"], [
        ("Allow", "Clean support text", "Approved fields continue to the local provider simulation"),
        ("Mask", "Email, phone, account reference", "Replace the identifier and preserve the support problem"),
        ("Block", "SSN, bank details, payment card, credentials, prompt manipulation", "Create no provider payload and keep the case in manual support"),
        ("Manual fallback", "Kill switch, timeout, provider error, invalid response, unknown category", "Stop safely and return the case to a support agent"),
    ], [1.25, 2.65, 2.85])
    doc.add_heading("Provider suppression in the code", level=1)
    add_paragraph(doc, "The full gateway calls the provider adapter only after preflight produces an approved payload. This check is the control that prevents a blocked ticket from reaching the provider path.")
    add_code(doc, '''preflight = prepare_provider_request(ticket, policy)
if preflight.provider_payload is None:
    return preflight

trace_steps = list(preflight.trace_steps)
raw_response = provider.invoke(preflight.provider_payload)''')
    doc.add_heading("Masking in the code", level=1)
    add_paragraph(doc, "The inspection layer stores the data type, action, and match position without copying the detected value into the finding record. It replaces only findings assigned the mask action.")
    add_code(doc, '''maskable = [finding for finding in findings
            if finding.action == FindingAction.MASK]
for finding in reversed(maskable):
    replacement = f"[MASKED_{finding.data_type.upper()}]"
    protected = protected[:finding.start] + replacement + protected[finding.end:]''')
    doc.add_heading("Technical safeguards", level=1)
    add_bullets(doc, [
        "Strict Pydantic models reject unexpected fields.",
        "Unicode and whitespace normalization reduce simple detection bypasses.",
        "A Luhn check helps distinguish payment-card candidates from arbitrary digit strings.",
        "The final gate examines the exact serialized outbound request.",
        "Response validation checks schema, category, size, markup, and sensitive content.",
        "Evidence receipts omit detected values and record decisions, policy versions, traceability, and limitations.",
    ])
    doc.add_heading("Demonstrated limits", level=1)
    add_paragraph(doc, "The provider in this local demonstration is an HTTPX MockTransport, so no real external AI service receives the ticket. The selected rules are deterministic and do not detect every possible format or manipulation. The application demonstrates a boundary-control pattern and tested behavior, not production readiness or universal detection accuracy.")
    return finish(doc, "03-python-security-gateway-technical-report.docx")


def validation_document():
    doc = Document(); style_doc(doc)
    add_title(doc, "SecureLaunch AI Validation Evidence", "Automated Python results and independent Amazon Bedrock Guardrail testing", "September 11-15, 2026")
    add_scope_note(doc)
    doc.add_heading("Executive result", level=1)
    add_paragraph(doc, "I validated the data-boundary concept in two separate environments. The local test suite verified the Python gateway. I then independently configured Amazon Bedrock Guardrails and tested managed sensitive-information and prompt-attack controls. The Streamlit demonstration does not call AWS, and the AWS lab did not invoke a foundation model.")
    add_table(doc, ["Evidence", "Observed result"], [
        ("Automated Python tests", "43 passed"),
        ("Local synthetic scenarios available in the interface", "15"),
        ("Selected AWS scenarios", "9 of 9 matched the expected outcome"),
        ("Configured AWS controls", "12 built-in sensitive-data types plus 1 custom account-reference regex"),
        ("Guardrail version", "Working draft tested, then Version 1 saved and smoke-tested"),
        ("Direct API reproduction", "ApplyGuardrail Version 1 masked phone and custom account reference"),
        ("Foundation-model calls in AWS lab", "0"),
        ("Real customer records", "0"),
    ], [3.0, 3.7])
    doc.add_heading("AWS decision logic", level=1)
    add_table(doc, ["Decision", "Purpose", "Examples"], [
        ("Allow", "Preserve legitimate support work when configured risks are absent", "Clean upload problem; harmless quoted instruction"),
        ("Mask", "Remove an unnecessary identifier while retaining useful context", "Email, phone number, LindenArc account reference"),
        ("Block", "Stop the request when restricted data or direct manipulation appears", "SSN, payment-card data, AWS credentials, prompt attack"),
    ], [1.0, 3.0, 2.7])
    doc.add_heading("Nine selected AWS tests", level=1)
    add_table(doc, ["ID", "Security question", "Observed result"], [
        ("T01", "Can a clean support ticket continue?", "Allowed; no Guardrail intervention"),
        ("T02", "Can an email be removed without discarding the ticket?", "Email masked"),
        ("T03", "Can managed and custom rules work together?", "Phone and account reference masked"),
        ("T04", "Does restricted identity data stop processing?", "SSN blocked"),
        ("T05", "Does payment-card data stop processing?", "Card number, expiry, and CVV blocked"),
        ("T06", "Do exposed cloud credentials stop processing?", "Access key and secret key blocked"),
        ("T07", "Does direct instruction manipulation stop processing?", "Prompt attack blocked with High confidence"),
        ("T08", "Can harmless quoted text remain usable?", "Allowed; no prompt attack detected"),
        ("T09", "Is simulated AI output inspected?", "Email in output masked"),
    ], [0.55, 3.45, 2.7])
    doc.add_heading("Custom account-reference rule", level=1)
    add_paragraph(doc, "I configured the pattern ACCT-(SYNTH-)?[0-9]{4,12} to recognize the fictional organization's account format. The rule requires the ACCT- prefix, permits the optional SYNTH- label used in the portfolio, and accepts 4 to 12 digits. The rule demonstrates why an organization may need controls beyond built-in PII detectors.")
    doc.add_heading("What this evidence supports", level=1)
    add_bullets(doc, [
        "Selected identifiers can be removed while preserving the support problem.",
        "Selected restricted identity, financial, credential, and manipulation cases can stop the request.",
        "A custom regex can extend managed detection to an organization-specific format.",
        "A fixed Guardrail version and direct API call improve reproducibility beyond console-only testing.",
    ])
    doc.add_heading("Evidence limitations", level=1)
    add_paragraph(doc, "These results describe a bounded synthetic sample. They do not measure production false-positive or false-negative rates, test every possible data format or prompt attack, validate a real external model, verify the external provider, or establish legal compliance. The AWS evidence strengthens the control case without replacing production security testing.")
    return finish(doc, "04-validation-evidence-report.docx")


def reflection_document():
    doc = Document(); style_doc(doc)
    add_title(doc, "SecureLaunch AI Risk Treatment and Lessons", "Security decisions, remaining risks, and what the evidence demonstrates", "September 9-16, 2026")
    add_scope_note(doc)
    doc.add_heading("Purpose and conclusion", level=1)
    add_paragraph(doc, "This project taught me to treat AI adoption as a change to the organization's data flow, trust boundaries, suppliers, and operating process. A passing filter test answers an important question, but it does not answer every security question. I used the evidence to separate controls I implemented and tested from risks that still require other safeguards.")
    doc.add_heading("Decisions I made", level=1)
    add_table(doc, ["Risk I identified", "My response", "What remains"], [
        ("Tickets contain more data than an AI summary needs", "Allow-list fields, remove metadata and attachments, mask selected identifiers, block restricted content, and inspect final JSON", "Unknown formats and new prompt attacks can evade selected detectors"),
        ("A plausible AI response can still be wrong", "Validate the response, restrict categories, retrieve human-authored playbooks, and retain agent review", "Schema validation does not establish factual accuracy or effective human oversight"),
        ("A provider can fail or return unsafe output", "Use timeout handling, response rejection, manual fallback, and a kill switch", "Local simulations do not prove real recovery times or manual support capacity"),
        ("A data filter cannot secure the whole system", "Record tenant, supplier, access, lifecycle, and operational risks separately", "Production tenant isolation, supplier terms, permissions, deletion, and monitoring require separate evidence"),
    ], [1.8, 2.9, 2.0])
    doc.add_heading("Framework and law connections", level=1)
    add_paragraph(doc, "I used selected framework and legal references to organize the security reasoning. These mappings show relevance; they are not a certification or a legal compliance determination.")
    add_table(doc, ["Reference", "Project connection"], [
        ("NIST CSF 2.0 ID.AM and ID.RA", "Inventory data and routes, identify and prioritize risk, and connect treatment to evidence"),
        ("NIST CSF 2.0 PR.DS", "Minimize, inspect, mask, block, and protect evidence within the tested scope"),
        ("NIST CSF 2.0 GV.SC", "Document supplier requirements and unresolved provider evidence"),
        ("NIST AI RMF Govern, Map, Measure, Manage", "Define accountability, understand the use case, test selected controls, and choose risk responses"),
        ("GDPR Articles 5(1)(c), 25, 28, and 32 where applicable", "Consider minimization, protection by design, processor obligations, and appropriate security"),
        ("FTC Safeguards Rule for covered institutions", "Consider risk assessment, safeguards, service-provider oversight, and testing"),
    ], [2.65, 4.05])
    doc.add_heading("What I learned", level=1)
    add_bullets(doc, [
        "Security analysis should begin before an external AI request exists, because the safest sensitive value is the one that never leaves the organization.",
        "Masking is useful only when the business context remains usable and the identifier is unnecessary for the stated purpose.",
        "Blocking needs a safe operational destination. In this design, the original ticket remains available to a support agent.",
        "Input controls are incomplete without final-payload inspection and return-path validation.",
        "Managed cloud controls can strengthen evidence, but they do not replace application controls, supplier review, or production testing.",
        "Test results must be reported with limits so decision-makers do not mistake a bounded demonstration for complete assurance.",
    ])
    doc.add_heading("My contribution", level=1)
    add_paragraph(doc, "I developed the Python security gateway and automated tests, completed the risk and data analysis, and independently configured and tested Amazon Bedrock Guardrails. I used AI-assisted development for the visual implementation of the reviewer-facing Streamlit interface. I defined the interface content, visual direction, and acceptance criteria so it would communicate my technical work clearly.")
    doc.add_heading("Next evidence I would request", level=1)
    add_bullets(doc, [
        "Tenant-scoped authorization and cross-tenant negative tests.",
        "Provider data-use, retention, deletion, subprocessor, incident, and change evidence.",
        "Production least-privilege access, monitoring, and rollback evidence.",
        "Detection evaluation across more formats, adversarial cases, and false-positive examples.",
        "Human-review quality and workload evidence from a controlled pilot.",
        "Deletion, backup-aging, restore, and evidence-integrity tests.",
    ])
    return finish(doc, "05-risk-treatment-and-project-lessons.docx")


if __name__ == "__main__":
    for path in [scope_document(), analysis_document(), build_document(), validation_document(), reflection_document()]:
        print(path)
